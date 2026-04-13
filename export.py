"""Veri dışa aktarma modülü — CSV ve Excel formatları.

Endüstriciler için: pandas, Minitab, Excel uyumlu çıktılar.
Türkçe karakter desteği için CSV'lerde UTF-8 BOM kullanılır.
"""
import csv
import io
import time
import db
from schedule_mgr import DAY_NAMES, schedule_mgr
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


UTF8_BOM = '﻿'

# Türkçe gün listesi (timestamp'ten ders zamanı çıkarmak için)
TR_DAYS_SHORT = ['Pzt', 'Sal', 'Çar', 'Per', 'Cum', 'Cmt', 'Paz']

# Bölge ID → Türkçe ad
ZONE_LABELS = {'on': 'Ön Bölge', 'orta': 'Orta Bölge', 'arka': 'Arka Bölge', '-': '-'}

# İşlem ID → Türkçe ad
ACTION_LABELS = {'giris': 'Giriş', 'cikis': 'Çıkış', 'yenileme': 'Yenileme'}


def _csv_bytes(rows, header):
    """Header + satırlardan UTF-8 BOM'lu CSV byte stream üret."""
    buf = io.StringIO()
    buf.write(UTF8_BOM)
    writer = csv.writer(buf, delimiter=';', quoting=csv.QUOTE_MINIMAL)
    writer.writerow(header)
    writer.writerows(rows)
    return buf.getvalue().encode('utf-8')


def attendance_csv(student_no=None):
    rows = db.get_attendance(limit=10000, student_no=student_no)
    data = []
    for r in rows:
        ts = time.localtime(r['timestamp'])
        data.append([
            time.strftime('%Y-%m-%d', ts),
            time.strftime('%H:%M:%S', ts),
            TR_DAYS_SHORT[ts.tm_wday],
            r['student_no'],
            ZONE_LABELS.get(r['zone'], r['zone']),
            ACTION_LABELS.get(r['action'], r['action']),
        ])
    return _csv_bytes(data, ['Tarih', 'Saat', 'Gün', 'Öğrenci No', 'Bölge', 'İşlem'])


def runs_csv():
    runs = db.list_energy_runs(limit=500)
    data = []
    for r in runs:
        data.append([
            r['id'],
            r['finished_at'],
            r['total_minutes'],
            r['energy_baseline'],
            r['energy_auto'],
            r['savings_kwh'],
            r['savings_pct'],
            r['cost_saved_tl'],
            r['carbon_saved_kg'],
        ])
    header = ['ID', 'Bitiş Tarihi', 'Toplam Dakika', 'Bazhat Enerji (kWh)',
              'Otomasyonlu (kWh)', 'Tasarruf (kWh)', 'Tasarruf %',
              'Maliyet Tasarrufu (TL)', 'Karbon Azaltımı (kg CO2)']
    return _csv_bytes(data, header)


def class_stats_csv():
    stats = db.class_attendance_stats(schedule_mgr.list_all())
    data = []
    for s in stats:
        data.append([
            s['class_id'],
            s['class_name'],
            DAY_NAMES[s['day']],
            s['time'],
            s['expected'],
            s['unique_students'],
            s['checkins'],
            s['occupancy_rate'],
            'Aktif' if s['active'] else 'İPTAL',
        ])
    header = ['Ders ID', 'Ders Adı', 'Gün', 'Saat',
              'Beklenen Öğrenci', 'Benzersiz Giriş', 'Toplam Olay',
              'Doluluk Oranı %', 'Durum']
    return _csv_bytes(data, header)


# --- Multi-sheet Excel

HEADER_FILL = PatternFill(start_color='1F2A40', end_color='1F2A40', fill_type='solid')
HEADER_FONT = Font(bold=True, color='38BDF8')


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'A2'


def _autosize(ws):
    for col in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0 for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(length + 2, 30)


def full_xlsx():
    """Tüm verileri tek bir Excel dosyasında, ayrı sayfalarda."""
    wb = Workbook()

    # Sayfa 1 — Özet
    ws = wb.active
    ws.title = 'Özet'
    stats = db.attendance_stats()
    runs = db.list_energy_runs(limit=500)
    total_savings = sum(r['savings_kwh'] for r in runs)
    total_cost    = sum(r['cost_saved_tl'] for r in runs)
    total_carbon  = sum(r['carbon_saved_kg'] for r in runs)

    ws['A1'] = 'Akıllı Sınıf Otomasyonu — Özet Rapor'
    ws['A1'].font = Font(bold=True, size=14, color='38BDF8')
    ws.merge_cells('A1:B1')
    ws['A3'] = 'Rapor Tarihi'
    ws['B3'] = time.strftime('%Y-%m-%d %H:%M:%S')
    ws['A5'] = 'Toplam Yoklama Olayı'
    ws['B5'] = stats['total_events']
    ws['A6'] = 'Benzersiz Öğrenci Sayısı'
    ws['B6'] = stats['unique_students']
    ws['A7'] = 'Bugünkü Giriş Sayısı'
    ws['B7'] = stats['today_checkins']
    ws['A9'] = 'Tamamlanan Simülasyon Sayısı'
    ws['B9'] = len(runs)
    ws['A10'] = 'Toplam Tasarruf (kWh)'
    ws['B10'] = round(total_savings, 2)
    ws['A11'] = 'Toplam Maliyet Tasarrufu (TL)'
    ws['B11'] = round(total_cost, 2)
    ws['A12'] = 'Toplam Karbon Azaltımı (kg CO2)'
    ws['B12'] = round(total_carbon, 2)
    for row in range(3, 13):
        ws.cell(row=row, column=1).font = Font(bold=True)
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 22

    # Sayfa 2 — Yoklama
    ws = wb.create_sheet('Yoklama Kayıtları')
    headers = ['Tarih', 'Saat', 'Gün', 'Öğrenci No', 'Bölge', 'İşlem']
    ws.append(headers)
    for r in db.get_attendance(limit=10000):
        ts = time.localtime(r['timestamp'])
        ws.append([
            time.strftime('%Y-%m-%d', ts),
            time.strftime('%H:%M:%S', ts),
            TR_DAYS_SHORT[ts.tm_wday],
            r['student_no'],
            ZONE_LABELS.get(r['zone'], r['zone']),
            ACTION_LABELS.get(r['action'], r['action']),
        ])
    _style_header(ws, len(headers))
    _autosize(ws)

    # Sayfa 3 — Sim Koşmaları
    ws = wb.create_sheet('Sim Koşmaları')
    headers = ['ID', 'Bitiş Tarihi', 'Toplam Dakika',
               'Bazhat (kWh)', 'Otomasyonlu (kWh)', 'Tasarruf (kWh)',
               'Tasarruf %', 'TL', 'CO2 (kg)']
    ws.append(headers)
    for r in runs:
        ws.append([r['id'], r['finished_at'], r['total_minutes'],
                   r['energy_baseline'], r['energy_auto'], r['savings_kwh'],
                   r['savings_pct'], r['cost_saved_tl'], r['carbon_saved_kg']])
    _style_header(ws, len(headers))
    _autosize(ws)

    # Sayfa 4 — Ders Bazlı
    ws = wb.create_sheet('Ders Bazlı Doluluk')
    headers = ['Ders ID', 'Ders Adı', 'Gün', 'Saat',
               'Beklenen', 'Benzersiz Giriş', 'Toplam Olay',
               'Doluluk %', 'Durum']
    ws.append(headers)
    for s in db.class_attendance_stats(schedule_mgr.list_all()):
        ws.append([s['class_id'], s['class_name'], DAY_NAMES[s['day']], s['time'],
                   s['expected'], s['unique_students'], s['checkins'],
                   s['occupancy_rate'], 'Aktif' if s['active'] else 'İPTAL'])
    _style_header(ws, len(headers))
    _autosize(ws)

    # Sayfa 5 — Sim Günlük Detay (her koşmanın günlük breakdown'ı)
    ws = wb.create_sheet('Sim Günlük Detay')
    headers = ['Koşma ID', 'Tarih', 'Gün', 'Bazhat (kWh)', 'Otomasyonlu (kWh)', 'Tasarruf (kWh)']
    ws.append(headers)
    for r in runs:
        for d in r.get('daily', []):
            ws.append([r['id'], r['finished_at'], d.get('day', ''),
                       d.get('baseline', 0), d.get('auto', 0), d.get('savings', 0)])
    _style_header(ws, len(headers))
    _autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
